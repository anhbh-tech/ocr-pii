"""Dựng xlsx: lưới cột và hàng trống quyết định bảng có lệch hay không."""
from __future__ import annotations

import re

import pytest
from openpyxl import Workbook

from ocr_lab.preview import xem_truoc


@pytest.fixture
def bang_co_o_gop(tmp_path):
    """Bảng tái hiện đúng hình dạng phiếu đề nghị thanh toán thật.

    Điểm mấu chốt: một ô gộp CHIỀU DỌC (A4:A5) rồi một hàng TRỐNG ngay dưới
    nó, đúng như tệp của khách.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Standard"
    ws["A1"] = "Kính gửi:"
    ws["F1"] = "Ngày:"
    ws["H1"] = "19/9/2025"
    ws["A4"] = "Email:"
    ws.merge_cells("A4:A5")     # gộp DỌC 2 hàng
    # hàng 6 để trống hoàn toàn
    ws["A7"] = "ĐỀ NGHỊ THANH TOÁN"
    ws["E8"] = "Số: V2025-904"
    ws.column_dimensions["A"].width = 6.0
    ws.column_dimensions["E"].width = 24.0
    f = tmp_path / "t.xlsx"
    wb.save(f)
    return f


def _hang(html_str: str) -> list[str]:
    return re.findall(r"<tr[^>]*>(.*?)</tr>", html_str, re.S)


def test_hang_trong_giua_bang_duoc_giu(bang_co_o_gop):
    """Bỏ hàng trống làm `rowspan` thò xuống hàng kế tiếp và đẩy lệch nó.

    Đây là nguyên nhân THẬT của bảng lệch, không phải chuyện thẩm mỹ: ô
    `A4:A5` phát ra `rowspan="2"`; nếu hàng 5 bị xoá thì ô ấy chiếm chỗ của
    hàng 6, và mọi ô hàng 6 bị đẩy sang phải một cột.
    """
    h = xem_truoc(bang_co_o_gop)["html"]
    hang = _hang(h)
    # 8 hàng dữ liệu: không hàng nào bị bỏ giữa chừng
    assert len(hang) >= 8, f"mất hàng: chỉ còn {len(hang)}"
    assert any("rTrong" in x for x in re.findall(r"<tr[^>]*>", h)), \
        "hàng trống phải được đánh dấu, không phải bị xoá"


def test_luoi_cot_khop_so_cot_that(bang_co_o_gop):
    """`<col>` phải đúng bằng `max_column`.

    Excel lưu `column_dimensions` cho cả cột trống ngoài vùng dữ liệu; đếm
    theo nó thì lưới rộng hơn số ô thật và bảng lại lệch theo hướng khác.
    """
    h = xem_truoc(bang_co_o_gop)["html"]
    n_col = len(re.findall(r"<col ", h))
    assert n_col == 8, f"sheet 8 cột nhưng phát {n_col} thẻ <col>"


def test_do_rong_lay_tu_excel(bang_co_o_gop):
    """Độ rộng người soạn đặt phải đi vào HTML, không bị thay bằng số mặc định."""
    h = xem_truoc(bang_co_o_gop)["html"]
    rong = [int(x) for x in re.findall(r'<col style="width:(\d+)px"', h)]
    assert rong[0] < rong[4], \
        f"cột A (6.0) phải hẹp hơn cột E (24.0), nhận {rong[0]} và {rong[4]}"


def test_o_gop_giu_nguyen_nhip(bang_co_o_gop):
    h = xem_truoc(bang_co_o_gop)["html"]
    assert 'rowspan="2"' in h, "ô gộp dọc A4:A5 phải thành rowspan"
